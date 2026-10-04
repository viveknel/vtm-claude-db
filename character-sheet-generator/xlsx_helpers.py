"""
xlsx_helpers.py — reusable openpyxl building blocks for WoD20-style
Excel character sheets.

This implements the visual and interaction conventions documented in
references/style_guide.md. Import it into a per-game build script and
call these functions with that game's specific data (labels, values,
dot maxima, colors); don't fork or hand-roll the styling logic per game
— that's exactly the repeated work this module exists to avoid.

Requires: openpyxl (pip install openpyxl --break-system-packages)

Typical usage sketch (see the bottom of this file for a runnable demo):

    from openpyxl import Workbook
    import xlsx_helpers as xh

    wb = Workbook()
    ws = wb.active
    ws.title = "1. Front Sheet"

    xh.style_title(ws, 1, 1, 11, "MY GAME: CHARACTER SHEET")
    xh.style_subtitle(ws, 2, 1, 11, "yellow cells are for your entries")
    xh.style_section_header(ws, 4, 1, 11, "ATTRIBUTES")
    xh.style_subheader(ws, 5, 1, 3, "PHYSICAL")
    xh.add_dot_trait(ws, 6, label_col=1, input_col=2, dot_col=3,
                      label="Strength", value=1, max_dots=10)
    xh.add_wide_entry(ws, 7, first_col=2, last_col=3, value="")  # a 2-cell Name: field
    ...
    xh.apply_standard_sheet_setup(ws, print_area="A1:K40")
    wb.save("out.xlsx")

Two things every rebuild tends to get wrong on the first pass, worth
internalizing before you write a line of per-game code:

1. **Every single-cell value is boxed, not just entry cells.** A plain
   black formula/derived-value cell (Budget, Remaining, a key output)
   gets the exact same thin box border as a yellow entry cell —
   style_formula() and style_key_output() already do this for you, but
   if you ever set a cell's value directly without going through one of
   this module's style_*/add_* functions, remember the border.
2. **A value that spans merged columns is never a closed box** — it's
   open-right "ruled" instead (see apply_span_border()), even for a
   short field like a 2-cell Name: entry. Use add_wide_entry() for any
   entry that needs more than one column, rather than merging cells and
   styling only the anchor — the rest of the span is left completely
   unstyled otherwise, which looks like a box that's open on three
   sides instead of one.
"""

from dataclasses import dataclass

from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

# ---------------------------------------------------------------------------
# Functional colors — fixed across every game. See style_guide.md's
# "Palette" table for what each one means. Don't change these per game;
# change Palette (below) instead, which covers the brand colors only.
# ---------------------------------------------------------------------------
ENTRY_FILL = "FFF9C4"
ENTRY_FONT = "1F4E8C"
LABEL_FONT = "2B2B2B"
NOTE_FONT = "6B6B6B"
FORMULA_FONT = "000000"
GOOD_FILL = "D9EAD3"
BAD_FILL = "F4CCCC"
WHITE = "FFFFFF"

THIN = Side(style="thin", color="B7B7B7")  # soft gray, NOT black — confirmed against the reference
MEDIUM = Side(style="medium", color="000000")  # fallback only; title/image-placeholder use _medium_border() instead
BOX_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOX_MEDIUM = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
RULED_LEFT = Border(left=THIN, top=THIN, bottom=THIN)  # leftmost cell of a ruled row
RULED_MID = Border(top=THIN, bottom=THIN)  # other cells of a ruled row


def _medium_border(color):
    """A medium box border in a given color. The reference colors its
    title cell's and image-placeholder's medium border with the
    palette's section_bg (the darker brand shade), not black — use this
    instead of a bare Border(Side('medium', '000000'))."""
    side = Side(style="medium", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


@dataclass
class Palette:
    """Brand colors for one game's workbook — the one thing this skill
    deliberately varies per game. Pick a trio per style_guide.md's
    'Adapting the accent color' section; don't reuse these literal
    Relentless-Age defaults for a different game."""

    title_bg: str = "7A1F2B"
    section_bg: str = "541520"
    subheader_bg: str = "C9A227"


DEFAULT_PALETTE = Palette()


# ---------------------------------------------------------------------------
# Header / band styling
# ---------------------------------------------------------------------------

def style_title(ws, row, first_col, last_col, text, palette=DEFAULT_PALETTE):
    """Row-1-style title bar: full-width merge, brand color, white bold
    16pt, medium border all around in the palette's section_bg color
    (confirmed against the reference — not plain black)."""
    _merge_and_set(ws, row, first_col, last_col, text)
    cell = ws.cell(row=row, column=first_col)
    cell.font = Font(name=FONT_NAME, size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=palette.title_bg)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    _apply_border_to_merge(ws, row, first_col, last_col, _medium_border(palette.section_bg))


def style_subtitle(ws, row, first_col, last_col, text):
    """Row-2-style subtitle/instructions: full-width merge, gray 11pt,
    no fill."""
    _merge_and_set(ws, row, first_col, last_col, text)
    cell = ws.cell(row=row, column=first_col)
    cell.font = Font(name=FONT_NAME, size=11, color=NOTE_FONT)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_section_header(ws, row, first_col, last_col, text, palette=DEFAULT_PALETTE):
    """Major section band, e.g. 'ATTRIBUTES', 'ADVANTAGES'."""
    _merge_and_set(ws, row, first_col, last_col, text)
    cell = ws.cell(row=row, column=first_col)
    cell.font = Font(name=FONT_NAME, size=12, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=palette.section_bg)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def style_subheader(ws, row, first_col, last_col, text, palette=DEFAULT_PALETTE):
    """Column-group band within a section, e.g. 'PHYSICAL', 'TALENTS'.
    Gets a thin box border like everything else with visible content —
    the reference never leaves a subheader band borderless."""
    _merge_and_set(ws, row, first_col, last_col, text)
    cell = ws.cell(row=row, column=first_col)
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=palette.subheader_bg)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    _apply_border_to_merge(ws, row, first_col, last_col, BOX_THIN)


def style_footnote(ws, row, first_col, last_col, text, size=9):
    """Small gray note/citation text, e.g. a rules-page citation."""
    _merge_and_set(ws, row, first_col, last_col, text)
    cell = ws.cell(row=row, column=first_col)
    cell.font = Font(name=FONT_NAME, size=size, color=NOTE_FONT)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


# ---------------------------------------------------------------------------
# Labels
# ---------------------------------------------------------------------------

def style_label(cell, text, bold=False, align="left"):
    """Plain row/field label — not an input, no fill, no border. Use
    bold=True for standalone field labels (Name:, Player:); leave False
    for the dozens of per-row trait labels (Strength, Alertness, ...)."""
    cell.value = text
    cell.font = Font(name=FONT_NAME, size=10, bold=bold, color=LABEL_FONT)
    cell.alignment = Alignment(horizontal=align, vertical="center")


# ---------------------------------------------------------------------------
# Entry cells and dot-tracks
# ---------------------------------------------------------------------------

def style_entry(cell, value=0, align="center", number_format=None):
    """A yellow cell the player types into. Works for numbers or text —
    pass value="" and align='left' for a text field."""
    cell.value = value
    cell.font = Font(name=FONT_NAME, size=10, color=ENTRY_FONT)
    cell.fill = PatternFill("solid", fgColor=ENTRY_FILL)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = BOX_THIN
    if number_format:
        cell.number_format = number_format


def style_dot_display(cell, formula):
    """A read-only cell rendering a dot-track from a formula (see
    dot_formula() below to build the formula string)."""
    cell.value = formula
    cell.font = Font(name=FONT_NAME, size=10, color=FORMULA_FONT)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = BOX_THIN


def dot_formula(input_ref, max_dots=10):
    """Build the standard dot-track formula string for an input cell
    reference, e.g. dot_formula('B11', 10) ->
    '=IF(B11="","",REPT("\u25cf",B11)&REPT("\u25cb",10-B11))'.

    Use max_dots=10 for Attributes/Abilities/most powers/most
    Backgrounds; use max_dots=5 for a Virtue-equivalent triad if this
    game uses a 5-dot scale for those (confirm — don't assume)."""
    return (
        f'=IF({input_ref}="","",REPT("\u25cf",{input_ref})'
        f'&REPT("\u25cb",{max_dots}-{input_ref}))'
    )


def add_dot_trait(ws, row, label_col, input_col, dot_col, label, value=0, max_dots=10):
    """Write one full trait row in one call: label, yellow entry cell,
    dot-track formula cell. Columns are 1-based ints. Returns the entry
    cell's A1-style reference string (e.g. 'B11') so the caller can use
    it in further formulas (sums, budget checks, etc.)."""
    style_label(ws.cell(row=row, column=label_col), label)
    input_cell = ws.cell(row=row, column=input_col)
    style_entry(input_cell, value)
    input_ref = f"{get_column_letter(input_col)}{row}"
    dot_cell = ws.cell(row=row, column=dot_col)
    style_dot_display(dot_cell, dot_formula(input_ref, max_dots))
    return input_ref


def style_formula(cell, formula, bold=False):
    """A plain black formula cell (subtotal, derived value) with no
    fill but a full thin-gray box border, same as an entry cell —
    confirmed against the reference, where every single-cell formula
    output (Budget, Extra Spent, Remaining, Total, etc.) is boxed
    exactly like a yellow entry cell; only genuinely plain label text
    stays borderless. bold=True for a subtotal/total row, False for an
    ordinary derived value. If this value spans a merged multi-column
    range, call apply_span_border() right after to replace this box
    with the open-right ruled style instead (see that function)."""
    cell.value = formula
    cell.font = Font(name=FONT_NAME, size=10, bold=bold, color=FORMULA_FONT)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BOX_THIN


def style_key_output(cell, formula):
    """Static (non-conditional) highlight for a number the player
    should transcribe elsewhere — starting Willpower, a running XP
    balance, a computed starting resource-pool total. Boxed like any
    other single-cell value (see style_formula); call
    apply_span_border() afterward if this spans a merged range."""
    cell.value = formula
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color=FORMULA_FONT)
    cell.fill = PatternFill("solid", fgColor=GOOD_FILL)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BOX_THIN


def apply_span_border(ws, row, first_col, last_col):
    """Apply the reference's 'ruled, open on the right' border to every
    cell in a merged value span on one row: a left wall on the first
    cell only, top+bottom on every cell, no right wall anywhere.

    The reference NEVER puts a closed box around a merged span, even a
    short two-cell span like a "Name:" entry field — it always uses
    this open-right ruled style instead, the same treatment documented
    below for free-text blocks, just generalized to any value (entry,
    formula, or key-output) that spans more than one merged column.

    This matters because merging cells and then styling only the
    anchor — e.g. `ws.merge_cells("B5:C5")` followed by
    `style_entry(ws.cell(row=5, column=2), ...)` — leaves every other
    cell in that merge completely unstyled: no border, no fill. Visually
    that reads as a box that's open on three sides instead of one,
    since only the anchor cell (B5) got a border at all. Call this
    right after merge_cells() + styling the anchor with style_entry() /
    style_formula() / style_key_output() to fix that — it overrides the
    anchor's box border with the ruled-left style and fills in proper
    borders on the rest of the span. For a one-line shortcut on a plain
    yellow entry, use add_wide_entry() instead of doing this by hand."""
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = RULED_LEFT if col == first_col else RULED_MID


def add_wide_entry(ws, row, first_col, last_col, value="", align="left"):
    """Yellow entry cell spanning a merged range on one row, in one
    call: merge + style_entry + apply_span_border. Use this any time an
    entry field needs more than one column — a 2-cell "Name:" box, a
    wide "Description:" field, a Background name slot — instead of
    merging and styling the anchor by hand and accidentally leaving the
    rest of the span unstyled (see apply_span_border's docstring)."""
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    style_entry(ws.cell(row=row, column=first_col), value, align=align)
    apply_span_border(ws, row, first_col, last_col)


# ---------------------------------------------------------------------------
# Free-text blocks and image placeholders
# ---------------------------------------------------------------------------

def add_free_text_block(ws, start_row, num_lines, first_col, last_col):
    """One ruled, enterable ('yellow') row per line — match num_lines to
    however many blank lines the source PDF gave this field. Returns
    the last row used."""
    for i in range(num_lines):
        row = start_row + i
        if first_col != last_col:
            ws.merge_cells(
                start_row=row, start_column=first_col,
                end_row=row, end_column=last_col,
            )
        for col in range(first_col, last_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = PatternFill("solid", fgColor=ENTRY_FILL)
            cell.font = Font(name=FONT_NAME, size=10, color=ENTRY_FONT)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = RULED_LEFT if col == first_col else RULED_MID
        ws.row_dimensions[row].height = 15
    return start_row + num_lines - 1


def add_image_placeholder(ws, first_row, last_row, first_col, last_col, text, palette=DEFAULT_PALETTE):
    """A boxed, non-entry block reserved for the player to paste in a
    picture (Insert > Picture) — not a formula-driven cell. The medium
    border is colored with the palette's section_bg (the darker brand
    shade), matching style_title — confirmed against the reference,
    not plain black."""
    ws.merge_cells(
        start_row=first_row, start_column=first_col,
        end_row=last_row, end_column=last_col,
    )
    cell = ws.cell(row=first_row, column=first_col)
    cell.value = text
    cell.font = Font(name=FONT_NAME, size=9, italic=True, color=NOTE_FONT)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _apply_border_to_merge(
        ws, first_row, first_col, last_col, _medium_border(palette.section_bg), last_row=last_row
    )


# ---------------------------------------------------------------------------
# Data validation and conditional formatting
# ---------------------------------------------------------------------------

def add_list_validation(ws, cell_range, options):
    """Attach a dropdown list to a cell or range, e.g.
    add_list_validation(ws, 'B8', ['Primary', 'Secondary', 'Tertiary']).
    """
    dv = DataValidation(
        type="list", formula1='"' + ",".join(options) + '"', allow_blank=True
    )
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv


def add_whole_number_validation(ws, cell_range, minimum=0, maximum=None):
    """Restrict a cell/range to whole numbers >= minimum (and <=
    maximum if given). Used on plain numeric entry cells that don't
    get a dot-track (Other Traits, table columns, etc.)."""
    dv = DataValidation(
        type="whole", operator="greaterThanOrEqual",
        formula1=str(minimum), allow_blank=True,
    )
    if maximum is not None:
        dv.operator = "between"
        dv.formula2 = str(maximum)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv


def add_budget_flag(ws, cell_ref, final_tally=False):
    """Conditional formatting for a 'Remaining' cell on the Character
    Creation sheet: red when it goes negative (overspent). Pass
    final_tally=True only for the sheet's single final is-chargen-done
    cell (usually 'Remaining Bonus/Freebie Points') to also add a green
    'resolved' rule for >= 0 — every other Remaining cell along the way
    only needs the red warning."""
    red = PatternFill("solid", bgColor=BAD_FILL)
    ws.conditional_formatting.add(
        cell_ref, CellIsRule(operator="lessThan", formula=["0"], fill=red)
    )
    if final_tally:
        green = PatternFill("solid", bgColor=GOOD_FILL)
        ws.conditional_formatting.add(
            cell_ref,
            CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=green),
        )


# ---------------------------------------------------------------------------
# Sheet-level setup
# ---------------------------------------------------------------------------

def apply_standard_sheet_setup(ws, print_area, freeze="A4"):
    """Apply the fixed view/print conventions to a finished sheet:
    gridlines off, freeze panes below the title block, landscape,
    fit-to-width-one-page (unconstrained height), explicit print area.
    Call this last, after all content/columns/rows are in place."""
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = print_area


def set_column_widths(ws, widths):
    """widths: dict of column letter -> width, e.g. {'A': 17, 'B': 6.51}."""
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def set_row_height(ws, row, height=15.0):
    ws.row_dimensions[row].height = height


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _merge_and_set(ws, row, first_col, last_col, text):
    if first_col != last_col:
        ws.merge_cells(
            start_row=row, start_column=first_col,
            end_row=row, end_column=last_col,
        )
    ws.cell(row=row, column=first_col).value = text


def _apply_border_to_merge(ws, first_row, first_col, last_col, border, last_row=None):
    last_row = last_row or first_row
    for r in range(first_row, last_row + 1):
        for c in range(first_col, last_col + 1):
            ws.cell(row=r, column=c).border = border


# ---------------------------------------------------------------------------
# Runnable smoke test / usage demo. Not part of the public API — run this
# file directly (`python3 xlsx_helpers.py`) to sanity-check the helpers
# produce a valid, openable workbook before relying on them for real work.
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "1. Demo Sheet"

    style_title(ws, 1, 1, 11, "DEMO GAME: CHARACTER SHEET")
    style_subtitle(ws, 2, 1, 11, "yellow cells are for your entries")
    style_section_header(ws, 4, 1, 11, "ATTRIBUTES (demo: 7/5/3 dots)")
    style_subheader(ws, 5, 1, 3, "PHYSICAL")
    style_subheader(ws, 5, 5, 7, "SOCIAL")
    refs = []
    refs.append(add_dot_trait(ws, 6, 1, 2, 3, "Strength", value=1))
    refs.append(add_dot_trait(ws, 7, 1, 2, 3, "Dexterity", value=1))
    refs.append(add_dot_trait(ws, 6, 5, 6, 7, "Charisma", value=1))
    style_formula(ws.cell(row=9, column=1), f"=SUM({refs[0]},{refs[1]})", bold=True)
    style_key_output(ws.cell(row=9, column=3), "=MAX(2,3)")

    add_free_text_block(ws, 11, 3, 1, 7)
    add_image_placeholder(ws, 15, 20, 1, 3, "(space reserved for a sketch)")

    add_list_validation(ws, "B25", ["Primary", "Secondary", "Tertiary"])
    style_entry(ws.cell(row=25, column=2), "Primary", align="center")
    style_formula(ws.cell(row=25, column=3), "=IF(B25=\"Primary\",7,IF(B25=\"Secondary\",5,3))")
    add_budget_flag(ws, "D25", final_tally=True)
    ws.cell(row=25, column=4, value=-1)

    set_column_widths(ws, {"A": 17, "B": 6.5, "C": 15, "D": 2.2,
                            "E": 17, "F": 6.5, "G": 15})
    apply_standard_sheet_setup(ws, print_area="A1:K30")

    out_path = sys.argv[1] if len(sys.argv) > 1 else "/tmp/xlsx_helpers_selftest.xlsx"
    wb.save(out_path)
    print(f"Self-test workbook written to {out_path}")

    # Re-open and do a couple of sanity checks.
    from openpyxl import load_workbook
    check = load_workbook(out_path, data_only=False)
    ws2 = check["1. Demo Sheet"]
    assert ws2["C6"].value.startswith("=IF(B6="), "dot formula missing"
    assert ws2["B6"].fill.fgColor.rgb == "00FFF9C4" or ws2["B6"].fill.fgColor.rgb == "FFFFF9C4", \
        "entry fill color wrong"
    assert len(ws2.conditional_formatting._cf_rules) >= 1, "conditional formatting missing"
    print("Self-test checks passed.")
