"""
Recalculation test for the V20 Dark Ages character sheet.

openpyxl never evaluates formulas, so a structurally clean workbook can still
miscalculate. This script fills copies of the workbook with test characters,
has LibreOffice recalculate them headlessly, and compares the results with
hand-derived expectations:

  * Magda - the book's own worked example (V20 Dark Ages ch. 4, printed pp.152-155):
    Road 7, Willpower 3, Eleventh Generation, max Blood Pool 12, starting Blood Pool 2,
    15 freebie points spent and 0 remaining.
  * Edge case A - Nosferatu, a bad priority assignment, Generation 5, Flaws over the
    cap, an overspent freebie budget, Road 10.
  * Edge case B - Generation 0 with a Blood Pool roll above the cap, Road 1.

Usage:
    python verify_recalc.py [path/to/workbook.xlsx]

With no argument it first builds a fresh workbook with build_v20_dark_ages_sheet.py and
tests that, so the generator is checked end to end. Needs openpyxl and LibreOffice (set
the SOFFICE environment variable if `soffice` is not on PATH). Exits non-zero if any
check fails. Never modifies the workbook under test.
"""
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent

STEP6_BUDGET = "Freebie budget (base + Flaw points up to the cap)"
STARTING_BP = "Starting Blood Pool (roll + Domain + Herd, capped)"
GENERATION_DOTS = "Generation Background dots (0 = Twelfth Generation)"


def find_soffice():
    candidates = [os.environ.get("SOFFICE"), shutil.which("soffice"),
                  r"C:\Program Files\LibreOffice\program\soffice.exe",
                  r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                  "/usr/bin/soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for c in candidates:
        if c and Path(c).exists():
            return c
    sys.exit("LibreOffice not found. Install it or set the SOFFICE environment variable.")


def find(ws, text, start=1, col=1):
    """Row of the first cell in `col` (from row `start`) whose value equals `text`."""
    for r in range(start, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == text:
            return r
    raise KeyError(f"label not found: {text!r}")


# ---------------------------------------------------------------------------
# Test inputs
# ---------------------------------------------------------------------------

def fill_magda(wb):
    cc, xp, fs = wb["5. Character Creation"], wb["6. Experience & Advancement"], wb["1. Front Sheet"]

    # Step 2: Mental primary, Physical secondary, Social tertiary
    cc.cell(row=find(cc, "Physical"), column=2).value = "Secondary"
    cc.cell(row=find(cc, "Social"), column=2).value = "Tertiary"
    cc.cell(row=find(cc, "Mental"), column=2).value = "Primary"
    a0 = find(cc, "Strength")
    for i, v in enumerate([3, 3, 2, 3, 2, 1, 3, 4, 3]):  # Str Dex Sta | Cha Man App | Per Int Wits
        cc.cell(row=a0 + i, column=3).value = v

    # Step 3: Knowledges primary (13), Talents secondary (8 of 9), Skills tertiary (5)
    for label, prio, dots in [("Knowledges", "Primary", 13), ("Talents", "Secondary", 8),
                              ("Skills", "Tertiary", 5)]:
        r = find(cc, label)
        cc.cell(row=r, column=2).value = prio
        cc.cell(row=r, column=4).value = dots

    # Step 4: Disciplines 2/1/1, Backgrounds 3/1/1, Virtues Conscience 3 / Self-Control 4 / Courage 3
    d0 = find(cc, "Discipline") + 1
    for i, (nm, v) in enumerate([("Dominate", 2), ("Fortitude", 1), ("Presence", 1)]):
        cc.cell(row=d0 + i, column=1).value = nm
        cc.cell(row=d0 + i, column=2).value = v
    b0 = find(cc, "Background") + 1
    for i, (nm, v) in enumerate([("Mentor", 3), ("Contacts", 1), ("Generation", 1)]):
        cc.cell(row=b0 + i, column=1).value = nm
        cc.cell(row=b0 + i, column=2).value = v
    v0 = find(cc, "Conscience / Conviction")
    for i, v in enumerate([3, 4, 3]):
        cc.cell(row=v0 + i, column=2).value = v

    # Step 5: Generation 1 dot (Eleventh), Blood Pool roll of 2, no Domain or Herd
    cc.cell(row=find(cc, GENERATION_DOTS), column=2).value = 1
    cc.cell(row=find(cc, "Starting Blood Pool roll (one die)"), column=2).value = 2

    # Step 6: two Attribute dots (10) + five Willpower dots (5) = the book's 15 points
    t0 = find(cc, "Trait")
    cc.cell(row=find(cc, "Attributes", start=t0), column=3).value = 2
    cc.cell(row=find(cc, "Willpower", start=t0), column=3).value = 5

    # Sheet 6: a FLAT, OLD, multi-step OLD and NEW purchase plus an overspend
    xp.cell(row=find(xp, "Starting XP (banked before you begin logging)"), column=4).value = 20
    hdr = find(xp, "Date / Session")
    rows = [("S1", 5, "Attribute", "Intelligence", 3, 4),            # 3 x 4         = 12
            ("S1", None, "Ability", "Occult", 2, 3),                 # 2 x 2         = 4
            ("S2", None, "New Discipline", "Auspex", 0, 1),          # flat          = 10
            ("S2", None, "Blood Sorcery Ritual", "Rite", None, 3),   # 3 x 2         = 6
            ("S3", None, "Willpower", None, 3, 5),                   # (3 + 4) x 1   = 7
            ("S3", None, "Area of Expertise", "Occult: Wards", None, None)]  # flat  = 1
    for i, row in enumerate(rows):
        for col, v in enumerate(row, 1):
            xp.cell(row=hdr + 1 + i, column=col).value = v

    fs["F38"].value = 7   # Road rating on the Front Sheet
    fs["B10"].value = 3   # dot-track spot checks
    fs["J30"].value = 3


def fill_edge_a(wb):
    cc, fs = wb["5. Character Creation"], wb["1. Front Sheet"]
    cc.cell(row=find(cc, "Is the character a Nosferatu?"), column=2).value = "Yes"
    cc.cell(row=find(cc, "Appearance"), column=3).value = 0
    cc.cell(row=find(cc, "Social"), column=2).value = "Primary"       # two Primaries, no Tertiary
    g = find(cc, GENERATION_DOTS)
    cc.cell(row=g, column=2).value = 5
    cc.cell(row=g + 5, column=2).value = 3                              # Domain
    cc.cell(row=g + 6, column=2).value = 3                              # Herd
    cc.cell(row=g + 7, column=2).value = 10                             # roll
    cc.cell(row=find(cc, "Flaw points taken"), column=2).value = 10     # cap is 7
    t0 = find(cc, "Trait")
    cc.cell(row=find(cc, "Attributes", start=t0), column=3).value = 4   # 20 points
    cc.cell(row=find(cc, "Willpower", start=t0), column=3).value = 4    # 4 points
    fs["F38"].value = 10


def fill_edge_b(wb):
    fill_edge_a(wb)
    cc = wb["5. Character Creation"]
    cc.cell(row=find(cc, GENERATION_DOTS), column=2).value = 0
    wb["1. Front Sheet"]["F38"].value = 1


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

class Checker:
    def __init__(self):
        self.total = 0
        self.failures = []

    def __call__(self, name, actual, expected):
        self.total += 1
        ok = actual == expected or (isinstance(actual, (int, float)) and isinstance(expected, (int, float))
                                    and abs(actual - expected) < 1e-9)
        if not ok:
            self.failures.append(f"{name}: got {actual!r}, expected {expected!r}")
        return ok


def check_magda(wb, check):
    cc, xp, fs = wb["5. Character Creation"], wb["6. Experience & Advancement"], wb["1. Front Sheet"]
    val = lambda ws, r, c: ws.cell(row=r, column=c).value
    blank = lambda v: v in (None, "")

    check("Front: Strength 3 dot-track", fs["C10"].value, "\u25cf\u25cf\u25cf\u25cb\u25cb\u25cb\u25cb\u25cb\u25cb\u25cb")
    check("Front: Alertness starts at 0", fs["C16"].value, "\u25cb" * 10)
    check("Front: blank custom Ability row stays blank", blank(fs["C26"].value), True)
    check("Front: blank Discipline row stays blank", blank(fs["C30"].value), True)
    check("Front: Virtue 5-dot track", fs["K30"].value, "\u25cf\u25cf\u25cf\u25cb\u25cb")
    check("Front: Virtue default 1", fs["K31"].value, "\u25cf\u25cb\u25cb\u25cb\u25cb")
    check("Front: aura modifier at Road 7", fs["F40"].value, "No modifier")

    for cat, (budget, free, spent, remaining) in {"Physical": (5, 3, 5, 0), "Social": (3, 3, 3, 0),
                                                  "Mental": (7, 3, 7, 0)}.items():
        r = find(cc, cat)
        check(f"Attributes {cat}: budget", val(cc, r, 3), budget)
        check(f"Attributes {cat}: free dots", val(cc, r, 4), free)
        check(f"Attributes {cat}: dots spent", val(cc, r, 5), spent)
        check(f"Attributes {cat}: remaining", val(cc, r, 6), remaining)
    pc = find(cc, "Priorities check")
    check("Attributes: priorities check", val(cc, pc, 2), "OK")

    for cat, (budget, rem) in {"Talents": (9, 1), "Skills": (5, 0), "Knowledges": (13, 0)}.items():
        r = find(cc, cat)
        check(f"Abilities {cat}: budget", val(cc, r, 3), budget)
        check(f"Abilities {cat}: remaining", val(cc, r, 5), rem)
    check("Abilities: priorities check", val(cc, find(cc, "Priorities check", start=pc + 1), 2), "OK")

    used = [r for r in range(1, cc.max_row + 1) if cc.cell(row=r, column=1).value == "Dots used"]
    for label, r, dots in zip(["Disciplines", "Backgrounds"], used, [4, 5]):
        check(f"{label}: dots used", val(cc, r, 2), dots)
        check(f"{label}: remaining", val(cc, r + 2, 2), 0)
    va = find(cc, "Dots above the free 3")
    check("Virtues: dots above free 3", val(cc, va, 2), 7)
    check("Virtues: remaining", val(cc, va + 2, 2), 0)

    road = find(cc, "Road rating (Cons./Conv. + Self-Control/Instinct)")
    check("Road rating (book: 7)", val(cc, road, 2), 7)
    check("Aura modifier at Road 7", val(cc, road + 1, 2), "No modifier")
    check("Starting Willpower (book: 3)", val(cc, road + 2, 2), 3)
    g = find(cc, GENERATION_DOTS)
    check("Generation (book: Eleventh)", val(cc, g + 1, 2), "Eleventh")
    check("Max Blood Pool (book: 12)", val(cc, g + 2, 2), 12)
    check("Blood points per turn at Eleventh", val(cc, g + 3, 2), 1)
    check("Max Trait dots at Eleventh", val(cc, g + 4, 2), 5)
    check("Starting Blood Pool (book: roll 2 -> 2)", val(cc, find(cc, STARTING_BP), 2), 2)

    t0 = find(cc, "Trait")
    check("Freebies: Attributes 2 x 5", val(cc, find(cc, "Attributes", start=t0), 4), 10)
    check("Freebies: Willpower 5 x 1", val(cc, find(cc, "Willpower", start=t0), 4), 5)
    check("Freebies spent (book: 15)", val(cc, find(cc, "Total freebie points spent"), 4), 15)
    check("Freebie budget, no Flaws", val(cc, find(cc, STEP6_BUDGET), 2), 15)
    check("Freebies remaining (book: 0)", val(cc, find(cc, "Remaining freebie points"), 2), 0)

    hdr = find(xp, "Date / Session")
    check("XP cost per log row", [val(xp, hdr + 1 + i, 7) for i in range(6)], [12, 4, 10, 6, 7, 1])
    check("XP running balance", [val(xp, hdr + 1 + i, 8) for i in range(6)], [13, 9, -1, -7, -14, -15])
    check("XP blank log row cost stays blank", blank(val(xp, hdr + 7, 7)), True)
    check("XP blank log row balance stays blank", blank(val(xp, hdr + 7, 8)), True)
    st = find(xp, "Starting XP (banked before you begin logging)")
    check("XP total earned", val(xp, st + 1, 4), 25)
    check("XP total spent", val(xp, st + 2, 4), 40)
    check("XP current balance", val(xp, st + 3, 4), -15)


def check_edge_a(wb, check):
    cc, fs = wb["5. Character Creation"], wb["1. Front Sheet"]
    val = lambda r, c: cc.cell(row=r, column=c).value
    soc = find(cc, "Social")
    check("A: Nosferatu Social free dots", val(soc, 4), 2)
    check("A: Nosferatu Social dots spent (3+2+0-2)", val(soc, 5), 3)
    check("A: Social as Primary budget", val(soc, 3), 7)
    check("A: Social remaining", val(soc, 6), 4)
    check("A: duplicate priorities flagged", val(find(cc, "Priorities check"), 2),
          "Assign Primary, Secondary and Tertiary exactly once")
    g = find(cc, GENERATION_DOTS)
    check("A: Generation at 5 dots", val(g + 1, 2), "Seventh")
    check("A: max Blood Pool at 5 dots", val(g + 2, 2), 20)
    check("A: blood per turn at 5 dots", val(g + 3, 2), 4)
    check("A: max Trait dots at 5 dots", val(g + 4, 2), 6)
    check("A: starting Blood Pool 10+3+3", val(find(cc, STARTING_BP), 2), 16)
    check("A: freebie budget 15 + min(10, cap 7)", val(find(cc, STEP6_BUDGET), 2), 22)
    check("A: freebies spent 20 + 4", val(find(cc, "Total freebie points spent"), 4), 24)
    check("A: freebies remaining", val(find(cc, "Remaining freebie points"), 2), -2)
    check("A: aura modifier at Road 10", fs["F40"].value, "-2 difficulty")


def check_edge_b(wb, check):
    cc, fs = wb["5. Character Creation"], wb["1. Front Sheet"]
    g = find(cc, GENERATION_DOTS)
    check("B: Generation at 0 dots", cc.cell(row=g + 1, column=2).value, "Twelfth")
    check("B: max Blood Pool at 0 dots", cc.cell(row=g + 2, column=2).value, 11)
    check("B: starting Blood Pool capped (16 -> 11)", cc.cell(row=find(cc, STARTING_BP), column=2).value, 11)
    check("B: aura modifier at Road 1", fs["F40"].value, "+2 difficulty")


# ---------------------------------------------------------------------------

def main():
    soffice = find_soffice()

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        if len(sys.argv) > 1:
            src = Path(sys.argv[1]).resolve()
            if not src.exists():
                sys.exit(f"workbook not found: {src}")
        else:
            src = tmp / "V20_Dark_Ages_Character_Sheet.xlsx"
            subprocess.run([sys.executable, str(HERE / "build_v20_dark_ages_sheet.py"), str(src)],
                           check=True, capture_output=True, text=True)
        inputs, recalced = tmp / "in", tmp / "recalc"
        inputs.mkdir()
        recalced.mkdir()
        variants = {"magda": [fill_magda], "edge_a": [fill_magda, fill_edge_a],
                    "edge_b": [fill_magda, fill_edge_b]}
        for name, fillers in variants.items():
            wb = load_workbook(src)
            for fill in fillers:
                fill(wb)
            wb.save(inputs / f"{name}.xlsx")

        # A list argument (not one string) keeps the multi-word filter name together.
        cmd = [soffice, "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
               "--outdir", str(recalced)] + [str(inputs / f"{n}.xlsx") for n in variants]
        subprocess.run(cmd, check=True, capture_output=True, text=True, timeout=300)

        check = Checker()
        for name, fn in [("magda", check_magda), ("edge_a", check_edge_a), ("edge_b", check_edge_b)]:
            out = recalced / f"{name}.xlsx"
            if not out.exists():
                sys.exit(f"LibreOffice did not produce {out.name}")
            fn(load_workbook(out, data_only=True), check)

    print(f"{check.total} checks run on {src.name}")
    if check.failures:
        print(f"{len(check.failures)} FAILED:")
        for f in check.failures:
            print("  -", f)
        sys.exit(1)
    print("all passed")


if __name__ == "__main__":
    main()
